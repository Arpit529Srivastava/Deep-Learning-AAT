"""Small helpers on top of python-docx that reproduce the DSCE AAT report layout."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor, Emu

FONT = "Times New Roman"
BROWN = "8B4513"


def set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)
    return run


def new_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6); pf.space_before = Pt(0); pf.line_spacing = 1.08
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(0.9); sec.bottom_margin = Inches(0.9)
    sec.header_distance = Inches(0.45); sec.footer_distance = Inches(0.4)
    return doc


def page_border(section, size=24):
    """Box border around the page (size in eighths of a point)."""
    sectPr = section._sectPr
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size)); b.set(qn("w:space"), "24"); b.set(qn("w:color"), "000000")
        pg.append(b)
    # pgBorders must come right after pgMar
    pgMar = sectPr.find(qn("w:pgMar"))
    pgMar.addnext(pg)


def para_border(p, side, sz=6, color="000000", space=1):
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr"); pPr.append(pbdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color)
    pbdr.append(b)


def add_field(run, instr):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = "1"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def setup_body_section(section, header_text, restart=False):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    set_font(hp.add_run(header_text), size=9.5)
    para_border(hp, "bottom", sz=4, color="808080", space=2)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    para_border(fp, "top", sz=24, color=BROWN, space=8)
    tab_pos = section.page_width - section.left_margin - section.right_margin
    fp.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
    set_font(fp.add_run("Dept. of ISE, DSCE\t\tPage "), size=9.5)
    r = fp.add_run(); set_font(r, size=9.5); add_field(r, "PAGE")
    if restart:
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
        section._sectPr.append(pg)


def para(doc, text="", size=11.5, bold=False, italic=False, align="justify", space_after=6, space_before=0, color=None):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(space_before)
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic, color=color)
    if size == 12.5 and align == "left":       # subheadings stay with what follows
        p.paragraph_format.keep_with_next = True
    return p


def add_rich(p, text, size=11.5, bold=False, italic=False, color=None):
    """Supports **bold** spans inside text."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            set_font(p.add_run(part), size=size, bold=bold or (i % 2 == 1), italic=italic, color=color)
    return p


def title(doc, text):
    p = para(doc, text, size=24, align="center", space_after=10, space_before=4)
    p.paragraph_format.keep_with_next = True
    return p


def heading(doc, text, level=1):
    size = 15 if level == 1 else 12.5
    p = para(doc, text, size=size, align="left", space_after=4, space_before=8 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    return p


def bullets(doc, items, numbered=False, size=11.5):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
        set_font(p.add_run(f"{i}.\t" if numbered else "•\t"), size=size)
        add_rich(p, it, size=size)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def cell_text(cell, text, bold=False, align="center", size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    add_rich(p, text, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, rows, widths=None, header=True, align=None, size=10.5, center=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER if center else WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            a = (align[j] if align else "center")
            cell_text(t.cell(i, j), str(val), bold=False, align=a, size=size)
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = Inches(w)
    # keep rows together and repeat header row
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i == 0 and header:
            th = OxmlElement("w:tblHeader"); trPr.append(th)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(2)
    return t


def figure(doc, path, caption, width=6.3):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))
    c = para(doc, caption, size=11, align="center", space_after=8)
    return c


def code_block(doc, code, size=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    tcPr = c._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    c.text = ""
    lines = code.strip("\n").split("\n")
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for k, line in enumerate(lines):
        r = p.add_run(line); set_font(r, size=size, name="Courier New")
        if k < len(lines) - 1:
            r.add_break()
    c.width = Inches(6.3)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(2)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_section(doc):
    return doc.add_section(WD_SECTION.NEW_PAGE)


def figure_pair(doc, left, right, caption, width=3.15):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, path in zip(t.rows[0].cells, (left, right)):
        c.width = Inches(width + 0.05)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run().add_picture(path, width=Inches(width))
    row = t.rows[0]._tr.get_or_add_trPr(); row.append(OxmlElement("w:cantSplit"))
    return para(doc, caption, size=11, align="center", space_after=8, space_before=3)
